"""Excel export of call results for the admin call-log page."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from agent.state import FIELD_LABELS

# call_outcome has its own column; everything else in FIELD_LABELS is a
# field captured on the call.
CAPTURED_FIELDS = [key for key in FIELD_LABELS if key != "call_outcome"]

HEADERS = (
    ["Batch", "Sheet Row", "Destination", "Extension", "Client", "Claim #", "DOI", "Patient", "Call Status"]
    + [FIELD_LABELS[key] for key in CAPTURED_FIELDS]
    + ["Started (UTC)", "Ended (UTC)", "Placed By", "Recording", "Notes", "Call SID"]
)


def _naive_utc(dt):
    # Excel has no timezone-aware datetime type; openpyxl refuses them.
    return dt.replace(tzinfo=None) if dt else None


def _recording_status(call_log) -> str:
    recordings = list(call_log.recordings.all())
    return recordings[0].status if recordings else ""


def build_call_log_workbook(call_logs) -> bytes:
    """One row per call. call_logs should have placed_by, batch and
    recordings loaded already."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Call Results"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for log in call_logs:
        claim = log.claim_context or {}
        captured = log.fields_captured or {}
        values = (
            [
                log.batch_id,
                log.row_number,
                log.destination_number,
                log.extension,
                claim.get("client_name", ""),
                claim.get("claim_number", ""),
                claim.get("doi", ""),
                claim.get("patient_name", ""),
                log.status_label,
            ]
            + [captured.get(key, "") for key in CAPTURED_FIELDS]
            + [
                # A queued or never-dialed row has a creation time, not a call time.
                _naive_utc(log.started_at) if log.call_sid else None,
                _naive_utc(log.ended_at),
                log.placed_by.username if log.placed_by else "",
                _recording_status(log),
                log.error,
                log.call_sid or "",
            ]
        )
        ws.append(values)
        # Captured values are whatever was said on the call. Force text so a
        # value that happens to start with "=" isn't run by Excel as a formula.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    for index, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = max(14, len(header) + 4)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
